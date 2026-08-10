#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
import_etf_screener.py — charge une fois Portefeuille_FNB.xlsx dans la table
`etf_screener` de TradeLab Pro, pour démarrer l'onglet ETF Screener avec les
fonds, les pondérations et les compositions déjà bâtis dans le classeur.

Usage :
  python tools/import_etf_screener.py [chemin_vers_Portefeuille_FNB.xlsx]

Sans argument : le classeur livré à côté de ce script (tools/Portefeuille_FNB.xlsx).

Idempotent : relancer ne duplique rien (upsert par ticker). En revanche, il
réécrit les colonnes présentes dans le classeur — si tu as modifié « Ma compo »
dans l'app depuis, la valeur du fichier Excel reprend le dessus. C'est un
import ponctuel, pas une synchronisation.

Prérequis : pip install openpyxl  (dépendance de cet outil seulement, pas de
l'application — rien dans TradeLab Pro n'ouvre de fichier Excel au démarrage).
"""
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Manque 'openpyxl'.  Fais :  pip install openpyxl")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from tradelab.data.database import Database  # noqa: E402

HEADER_ROW = 7            # ligne des en-têtes dans le classeur
DEFAULT_WORKBOOK = Path(__file__).resolve().parent / "Portefeuille_FNB.xlsx"

# En-tête du classeur -> colonne de la table etf_screener.
COLUMN_MAP = {
    "Ticker": "ticker",
    "Nom": "name",
    "Catégorie": "category",
    "Bourse": "exchange",
    "Devise": "currency",
    "Région": "region",
    "% Can": "pct_can",
    "% US": "pct_us",
    "% Intl": "pct_intl",
    "% Obl.": "pct_bond",
    "% Or/Alt": "pct_gold",
    "Risque": "risk",
    "MER": "mer",
    "Rend 1M": "ret_1m",
    "Rend 3M": "ret_3m",
    "Rend 6M": "ret_6m",
    "Rend 1A": "ret_1a",
    "Rend 3A": "ret_3a",
    "Rend 5A": "ret_5a",
    "Rend 10A": "ret_10a",
    "Volatilité (ann.)": "volatility",
    "Pire baisse": "max_drawdown",
    "Sharpe": "sharpe",
    # Le classeur nommait ses compositions Reco / Spéculatif ; l'onglet les
    # nomme par leur niveau de risque. Rien de « Low risk » dans le classeur :
    # cette colonne se remplit dans l'app.
    "Reco ★": "mid_risk",
    "Spéculatif ★": "high_risk",
    "Ma compo ✏️": "my_mix",
    "Notes": "notes",
    "Compte suggéré (REER/CELI)": "suggested_account",
    "Yahoo": "yahoo",
}

TEXT_COLUMNS = {"name", "category", "exchange", "currency", "region",
                "notes", "suggested_account", "yahoo"}


def read_header(ws) -> dict:
    """Position de chaque en-tête connu, par son libellé exact."""
    headers = {}
    for c in range(1, ws.max_column + 1):
        value = ws.cell(HEADER_ROW, c).value
        if value:
            headers[str(value).strip()] = c
    return headers


def read_rows(ws) -> list[dict]:
    """Un dict par fonds, prêt pour Database.etf_upsert(). Les cellules vides
    sont omises plutôt qu'écrites à zéro : une case laissée blanche dans le
    classeur veut dire « pas la donnée », pas « zéro pour cent »."""
    headers = read_header(ws)
    if "Ticker" not in headers:
        raise SystemExit(f"Colonne 'Ticker' introuvable (ligne d'en-tête = {HEADER_ROW} ?).")
    mapped = {col: db_col for label, db_col in COLUMN_MAP.items()
              if (col := headers.get(label)) is not None}

    funds = []
    r = HEADER_ROW + 1
    while ws.cell(r, headers["Ticker"]).value:
        fields = {}
        for col, db_col in mapped.items():
            value = ws.cell(r, col).value
            if value is None or value == "":
                continue
            if db_col in TEXT_COLUMNS:
                value = str(value).strip()
            fields[db_col] = value
        ticker = str(fields.pop("ticker", "")).strip().upper()
        if ticker:
            fields["ticker"] = ticker
            funds.append(fields)
        r += 1
    return funds


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    if not path.exists():
        sys.exit(f"Classeur introuvable : {path}")
    print(f"Fichier : {path}")

    # data_only=True : on veut les valeurs, pas les formules du bloc de totaux.
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Portefeuille"] if "Portefeuille" in wb.sheetnames else wb.active
    funds = read_rows(ws)

    db = Database()
    for fund in funds:
        ticker = fund.pop("ticker")
        db.etf_upsert(ticker, **fund)
    print(f"Terminé : {len(funds)} fonds importés dans etf_screener.")
    print(f"Base : {db.path}")
    print("Ouvre l'onglet « ETF Screener » puis « Rafraîchir rendements/risque » "
          "pour recalculer les rendements à partir de Yahoo Finance.")


if __name__ == "__main__":
    main()
