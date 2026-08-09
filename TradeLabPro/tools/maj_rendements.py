#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
maj_rendements.py  —  « Bouton » de mise à jour des rendements pour Portefeuille_FNB.xlsx

Ce que ça fait :
  Va chercher les prix récents (yfinance / Yahoo Finance) et réécrit les colonnes
  de rendement 1M / 3M / 6M / 1A / 3A / 5A / 10A pour chaque fonds du fichier Excel.
  Les rendements >= 1 an sont ANNUALISÉS. Dividendes réinvestis (cours ajustés).

Prérequis (une seule fois) :
  pip install yfinance openpyxl pandas

Utilisation :
  python maj_rendements.py                     (utilise Portefeuille_FNB.xlsx dans le dossier courant)
  python maj_rendements.py mon_fichier.xlsx    (autre fichier)

Notes :
  - Les fonds de moins de 10 ans : la colonne 10A n'est PAS écrasée (on garde la valeur approx. déjà en place).
  - Ouvre ensuite le fichier dans Excel : les totaux du bas se recalculent automatiquement.
  - Le script lit le symbole Yahoo dans la colonne « Yahoo » du fichier (TSX = suffixe .TO).
"""

import sys
import time
import datetime as dt

try:
    import openpyxl
except ImportError:
    sys.exit("Manque 'openpyxl'.  Fais :  pip install openpyxl")

try:
    import pandas as pd
except ImportError:
    sys.exit("Manque 'pandas'.  Fais :  pip install pandas")

try:
    import yfinance as yf
except ImportError:
    sys.exit("Manque 'yfinance'.  Fais :  pip install yfinance")


HEADER_ROW = 7          # ligne des en-têtes dans le fichier
SLEEP = 0.6             # pause entre requêtes (évite d'être bloqué par Yahoo)
RISK_FREE = 0.03        # taux sans risque annuel pour le ratio de Sharpe (ajustable)

# En-tête de colonne -> (nombre de jours en arrière, annualiser sur N années ou None)
PERIODS = {
    "Rend 1M":  (30,        None),
    "Rend 3M":  (91,        None),
    "Rend 6M":  (182,       None),
    "Rend 1A":  (365,       1),
    "Rend 3A":  (3 * 365,   3),
    "Rend 5A":  (5 * 365,   5),
    "Rend 10A": (10 * 365,  10),
}


def trailing_return(prices: "pd.Series", days: int, annualize_years):
    """Rendement total sur 'days' jours ; annualisé si annualize_years fourni. None si historique insuffisant."""
    if prices is None or prices.empty:
        return None
    end_date = prices.index[-1]
    end_price = float(prices.iloc[-1])
    target = end_date - pd.Timedelta(days=days)
    # Historique trop court ? (tolérance d'une semaine)
    if prices.index[0] > target + pd.Timedelta(days=7):
        return None
    sub = prices[prices.index <= target]
    if sub.empty:
        return None
    start_price = float(sub.iloc[-1])
    if start_price <= 0:
        return None
    total = end_price / start_price - 1.0
    if annualize_years:
        return (1.0 + total) ** (1.0 / annualize_years) - 1.0
    return total


def risk_metrics(prices: "pd.Series"):
    """Renvoie (volatilité annualisée, pire baisse, Sharpe) sur l'historique disponible. None si insuffisant."""
    if prices is None or len(prices) < 60:
        return None, None, None
    daily = prices.pct_change().dropna()
    if daily.empty:
        return None, None, None
    vol = float(daily.std()) * (252 ** 0.5)                       # volatilité annualisée
    running_max = prices.cummax()
    drawdown = prices / running_max - 1.0
    max_dd = float(drawdown.min())                                # pire baisse (négatif)
    years = (prices.index[-1] - prices.index[0]).days / 365.25
    if years <= 0 or float(prices.iloc[0]) <= 0:
        return round(vol, 4), round(max_dd, 4), None
    ann_ret = (float(prices.iloc[-1]) / float(prices.iloc[0])) ** (1.0 / years) - 1.0
    sharpe = (ann_ret - RISK_FREE) / vol if vol > 0 else None
    return round(vol, 4), round(max_dd, 4), (round(sharpe, 2) if sharpe is not None else None)


def get_prices(symbol: str):
    """Série de cours ajustés (Close) sur ~11 ans. None si échec."""
    try:
        df = yf.download(symbol, period="11y", interval="1d",
                         auto_adjust=True, progress=False, threads=False)
    except Exception as e:
        print(f"    ! téléchargement échoué ({e})")
        return None
    if df is None or df.empty:
        return None
    # yfinance peut renvoyer des colonnes multi-index selon la version
    if isinstance(df.columns, pd.MultiIndex):
        try:
            close = df["Close"][symbol]
        except Exception:
            close = df["Close"].iloc[:, 0]
    else:
        close = df["Close"]
    close = close.dropna()
    close.index = pd.to_datetime(close.index)
    return close.sort_index()


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else "Portefeuille_FNB.xlsx"
    print(f"Fichier : {path}")
    wb = openpyxl.load_workbook(path)
    ws = wb["Portefeuille"] if "Portefeuille" in wb.sheetnames else wb.active

    # Repérer les colonnes par leur en-tête
    headers = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(HEADER_ROW, c).value
        if v:
            headers[str(v).strip()] = c
    col_ticker = headers.get("Ticker")
    col_yahoo = headers.get("Yahoo")
    if not col_ticker or not col_yahoo:
        sys.exit("Colonnes 'Ticker' ou 'Yahoo' introuvables (ligne d'en-tête = 7 ?).")
    period_cols = {name: headers[name] for name in PERIODS if name in headers}
    col_vol = headers.get("Volatilité (ann.)")
    col_dd = headers.get("Pire baisse")
    col_sharpe = headers.get("Sharpe")

    updated, skipped = 0, 0
    r = HEADER_ROW + 1
    while ws.cell(r, col_ticker).value:
        ticker = str(ws.cell(r, col_ticker).value).strip()
        symbol = str(ws.cell(r, col_yahoo).value).strip()
        print(f"  {ticker:6} ({symbol}) ...", end="", flush=True)
        prices = get_prices(symbol)
        if prices is None or len(prices) < 20:
            print(" pas de données, ignoré")
            skipped += 1
            r += 1
            time.sleep(SLEEP)
            continue
        n_written = 0
        for name, col in period_cols.items():
            days, ann = PERIODS[name]
            val = trailing_return(prices, days, ann)
            if val is not None:
                cell = ws.cell(r, col, round(val, 4))   # stocké en fraction (0.25 -> 25,0%)
                cell.number_format = "0.0%"
                n_written += 1
            # si None (historique trop court) : on laisse la cellule telle quelle
        # Mesures de risque
        vol, max_dd, sharpe = risk_metrics(prices)
        if col_vol and vol is not None:
            c = ws.cell(r, col_vol, vol); c.number_format = "0.0%"
        if col_dd and max_dd is not None:
            c = ws.cell(r, col_dd, max_dd); c.number_format = "0.0%"
        if col_sharpe and sharpe is not None:
            c = ws.cell(r, col_sharpe, sharpe); c.number_format = "0.00"
        print(f" {n_written} période(s)")
        updated += 1
        r += 1
        time.sleep(SLEEP)

    out = path
    wb.save(out)
    print(f"\nTerminé : {updated} fonds mis à jour, {skipped} ignorés.")
    print(f"Enregistré dans {out}. Ouvre-le dans Excel pour recalculer les totaux du bas.")
    print(f"Date de rafraîchissement : {dt.date.today().isoformat()}")


if __name__ == "__main__":
    main()
